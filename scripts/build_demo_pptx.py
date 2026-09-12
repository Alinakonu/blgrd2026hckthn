#!/usr/bin/env python3
"""Build the AgriSense demo deck as a 16:9 PPTX matching the HTML theme."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "AgriSense-demo.pptx"
OUT_DOCS = ROOT / "docs" / "AgriSense-demo.pptx"

# Field-day tokens from the Streamlit / HTML deck
BG = RGBColor(0xE6, 0xED, 0xE7)
INK = RGBColor(0x1B, 0x2D, 0x24)
INK_SOFT = RGBColor(0x3E, 0x53, 0x49)
MUTED = RGBColor(0x6A, 0x7C, 0x72)
ACCENT = RGBColor(0x2F, 0x5F, 0x46)
DANGER = RGBColor(0xA8, 0x48, 0x40)
PANEL = RGBColor(0xFF, 0xFC, 0xF8)
PANEL_DANGER = RGBColor(0xFF, 0xF6, 0xF2)
PANEL_GOOD = RGBColor(0xEC, 0xF6, 0xEF)
RULE = RGBColor(0xC9, 0xD3, 0xCC)
SERIF = "Georgia"
SANS = "Calibri"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_run(run, text, size, color, font=SERIF, bold=False, tracking=None):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = font
    run.font.bold = bold
    if tracking is not None:
        rPr = run._r.get_or_add_rPr()
        rPr.set("spc", str(int(tracking * 100)))


def textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    return box, tf


def para(tf, text, size, color, font=SERIF, bold=False, align=PP_ALIGN.LEFT,
         space_before=0, space_after=0, tracking=None, first=False):
    p = tf.paragraphs[0] if first and tf.paragraphs else tf.add_paragraph()
    if first:
        # after clear(), there is one empty paragraph
        p = tf.paragraphs[0]
        p.clear()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    set_run(run, text, size, color, font, bold, tracking)
    return p


def mixed_para(tf, parts, align=PP_ALIGN.LEFT, space_before=0, space_after=0, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    if first:
        p = tf.paragraphs[0]
        p.clear()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    for text, size, color, font, bold in parts:
        run = p.add_run()
        set_run(run, text, size, color, font, bold)
    return p


def card(slide, left, top, width, height, fill=PANEL, radius=0.16):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = RULE
    sh.line.width = Pt(1)
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    return sh


def oval(slide, left, top, width, height, fill, alpha=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    if alpha is not None:
        spPr = sh.fill._xPr
        srgb = spPr.find(qn("a:solidFill"))
        if srgb is not None:
            srgbClr = srgb.find(qn("a:srgbClr"))
            if srgbClr is not None:
                etree.SubElement(srgbClr, qn("a:alpha"), val=str(int(alpha * 1000)))
    return sh


def set_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    # Soft sage blobs — same idea as the HTML atmosphere
    oval(slide, Inches(-1.8), Inches(-2.2), Inches(7.2), Inches(4.8), RGBColor(0xB4, 0xD2, 0xBE), alpha=55)
    oval(slide, Inches(8.6), Inches(-2.0), Inches(6.4), Inches(4.4), RGBColor(0xE8, 0xDC, 0xC8), alpha=35)
    oval(slide, Inches(3.2), Inches(5.4), Inches(8.0), Inches(4.2), RGBColor(0xA8, 0xC4, 0xB0), alpha=40)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    return slide


def eyebrow_box(slide, left, top, width, text, align=PP_ALIGN.LEFT):
    box, tf = textbox(slide, left, top, width, Inches(0.4))
    para(tf, text.upper(), 12, MUTED, SANS, True, align, tracking=14, first=True)
    return box


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1. Title
    s = blank(prs)
    eyebrow_box(s, Inches(1.2), Inches(2.05), Inches(10.9), "Serbia  ·  plot by plot", PP_ALIGN.CENTER)
    box, tf = textbox(s, Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.8))
    para(tf, "AgriSense", 80, INK, SERIF, True, PP_ALIGN.CENTER, first=True)
    box, tf = textbox(s, Inches(0.8), Inches(4.3), Inches(11.7), Inches(0.7))
    para(tf, "Climate-adaptive land copilot", 24, INK_SOFT, SANS, False, PP_ALIGN.CENTER, first=True)
    notes(s, "Imagine you farm around Novi Sad. Annual rainfall is basically the same. That number is lying. AgriSense: plot-level copilot for Serbia.")

    # 2. Questions
    s = blank(prs)
    eyebrow_box(s, Inches(1.7), Inches(1.55), Inches(10), "This plot")
    qs = [
        ("01", "What hit this field?"),
        ("02", "Is summer drying?"),
        ("03", "Keep maize — or change?"),
    ]
    y = 2.15
    for n, q in qs:
        box, tf = textbox(s, Inches(1.7), Inches(y), Inches(1.1), Inches(0.7))
        para(tf, n, 20, ACCENT, SERIF, True, first=True)
        box, tf = textbox(s, Inches(2.9), Inches(y - 0.08), Inches(9), Inches(0.8))
        para(tf, q, 36, INK, SERIF, True, first=True)
        y += 1.05
    notes(s, "Three questions a landowner actually needs answered — for this field, not the country.")

    # 3. Novi Sad lie
    s = blank(prs)
    eyebrow_box(s, Inches(1.35), Inches(1.45), Inches(10), "Novi Sad")
    left, top, w, h, gap = Inches(1.35), Inches(2.05), Inches(5.05), Inches(3.15), Inches(0.4)
    card(s, left, top, w, h, PANEL)
    box, tf = textbox(s, left + Inches(0.4), top + Inches(0.35), w - Inches(0.7), Inches(2.5))
    para(tf, "ANNUAL RAIN", 12, MUTED, SANS, True, tracking=12, first=True)
    mixed_para(tf, [
        ("599 ", 44, INK, SERIF, True),
        ("→ ", 36, MUTED, SERIF, False),
        ("609", 44, INK, SERIF, True),
    ], space_before=10)
    para(tf, "Looks fine.", 18, INK_SOFT, SANS, False, space_before=10)

    card(s, left + w + gap, top, w, h, PANEL_DANGER)
    box, tf = textbox(s, left + w + gap + Inches(0.4), top + Inches(0.35), w - Inches(0.7), Inches(2.5))
    para(tf, "SUMMER WATER", 12, MUTED, SANS, True, tracking=12, first=True)
    mixed_para(tf, [
        ("−241 ", 44, DANGER, SERIF, True),
        ("→ ", 36, MUTED, SERIF, False),
        ("−298", 44, DANGER, SERIF, True),
    ], space_before=10)
    para(tf, "−57 mm. Hidden.", 18, INK_SOFT, SANS, False, space_before=10)
    notes(s, "Novi Sad annual rain 599 to 609. Looks fine. Then point at summer water: minus 241 to minus 298. Fifty-seven millimetres more stress, hidden inside a stable year.")

    # 4. Maize vs wheat
    s = blank(prs)
    eyebrow_box(s, Inches(1.35), Inches(1.45), Inches(10), "Novi Sad  ·  same climate")
    card(s, left, top, w, h, PANEL_DANGER)
    box, tf = textbox(s, left + Inches(0.4), top + Inches(0.35), w - Inches(0.7), Inches(2.6))
    para(tf, "MAIZE", 12, MUTED, SANS, True, tracking=12, first=True)
    para(tf, "2.48", 52, DANGER, SERIF, True, space_before=6)
    para(tf, "Jun–Aug. 501k ha.", 18, INK_SOFT, SANS, False, space_before=8)
    pill = card(s, left + Inches(0.4), top + Inches(2.35), Inches(2.15), Inches(0.38), RGBColor(0xF8, 0xE8, 0xE6), radius=0.5)
    box, tf = textbox(s, left + Inches(0.4), top + Inches(2.36), Inches(2.15), Inches(0.36))
    para(tf, "highest exposure", 11, DANGER, SANS, True, PP_ALIGN.CENTER, first=True)

    card(s, left + w + gap, top, w, h, PANEL)
    box, tf = textbox(s, left + w + gap + Inches(0.4), top + Inches(0.35), w - Inches(0.7), Inches(2.6))
    para(tf, "WHEAT", 12, MUTED, SANS, True, tracking=12, first=True)
    para(tf, "stable", 52, INK, SERIF, True, space_before=6)
    para(tf, "Harvested first.", 18, INK_SOFT, SANS, False, space_before=8)
    card(s, left + w + gap + Inches(0.4), top + Inches(2.35), Inches(0.85), Inches(0.38), RGBColor(0xE4, 0xF0, 0xE8), radius=0.5)
    box, tf = textbox(s, left + w + gap + Inches(0.4), top + Inches(2.36), Inches(0.85), Inches(0.36))
    para(tf, "0.41", 11, ACCENT, SANS, True, PP_ALIGN.CENTER, first=True)
    notes(s, "Maize peak demand is June–August — the window that deteriorated. Irrigation gap 327 to 389. Exposure 2.48. Wheat is largely stable: off the field by early July.")

    # 5. Zrenjanin
    s = blank(prs)
    eyebrow_box(s, Inches(0.8), Inches(1.85), Inches(11.7), "Zrenjanin  ·  days ≥ 35°C", PP_ALIGN.CENTER)
    box, tf = textbox(s, Inches(0.5), Inches(2.35), Inches(12.3), Inches(2.4))
    mixed_para(tf, [
        ("2.7  ", 96, MUTED, SERIF, True),
        ("→  ", 72, MUTED, SERIF, False),
        ("8.0", 96, DANGER, SERIF, True),
    ], PP_ALIGN.CENTER, first=True)
    box, tf = textbox(s, Inches(0.8), Inches(4.85), Inches(11.7), Inches(0.6))
    para(tf, "Nearly triple.", 24, INK_SOFT, SANS, False, PP_ALIGN.CENTER, first=True)
    notes(s, "Zrenjanin: days at or above 35°C from 2.7 to 8.0. Almost a tripling. Summer deficit minus 260 to minus 312. Outlook 46 to 55 days above 30 — a range, not a forecast.")

    # 6. Niš
    s = blank(prs)
    eyebrow_box(s, Inches(1.35), Inches(0.85), Inches(10), "Niš")
    box, tf = textbox(s, Inches(1.35), Inches(1.2), Inches(10), Inches(1.7))
    para(tf, "No change", 48, INK, SERIF, True, first=True)
    para(tf, "indicated.", 48, INK, SERIF, True)
    cw, ch = Inches(5.05), Inches(2.15)
    cy = Inches(3.55)
    card(s, left, cy, cw, ch, PANEL_GOOD)
    box, tf = textbox(s, left + Inches(0.4), cy + Inches(0.3), cw - Inches(0.7), Inches(1.6))
    para(tf, "ANNUAL RAIN", 12, MUTED, SANS, True, tracking=12, first=True)
    mixed_para(tf, [
        ("521 ", 40, ACCENT, SERIF, True),
        ("→ ", 32, MUTED, SERIF, False),
        ("682", 40, ACCENT, SERIF, True),
    ], space_before=8)
    card(s, left + cw + gap, cy, cw, ch, PANEL_GOOD)
    box, tf = textbox(s, left + cw + gap + Inches(0.4), cy + Inches(0.3), cw - Inches(0.7), Inches(1.6))
    para(tf, "SUMMER WATER", 12, MUTED, SANS, True, tracking=12, first=True)
    mixed_para(tf, [
        ("−287 ", 40, ACCENT, SERIF, True),
        ("→ ", 32, MUTED, SERIF, False),
        ("−270", 40, ACCENT, SERIF, True),
    ], space_before=8)
    notes(s, "If the tool shouted crisis on every pin, you should not believe it. Niš annual rain 521 to 682. Summer deficit eased. Verdict: no change indicated.")

    # 7. How it works
    s = blank(prs)
    eyebrow_box(s, Inches(1.15), Inches(1.7), Inches(10), "Under the hood")
    items = [
        ("OPEN DATA", "Memory", "ERA5"),
        ("THIS CROP", "Exposure", "FAO-56"),
        ("THIS SEASON", "Plan", "Grok"),
    ]
    cw = Inches(3.5)
    gap3 = Inches(0.28)
    x0 = Inches(1.15)
    top7 = Inches(2.3)
    h7 = Inches(3.0)
    for i, (lab, title, hint) in enumerate(items):
        x = x0 + i * (cw + gap3)
        card(s, x, top7, cw, h7, PANEL)
        box, tf = textbox(s, x + Inches(0.35), top7 + Inches(0.4), cw - Inches(0.6), Inches(2.3))
        para(tf, lab, 12, MUTED, SANS, True, tracking=12, first=True)
        para(tf, title, 32, INK, SERIF, True, space_before=10)
        para(tf, hint, 18, INK_SOFT, SANS, False, space_before=10)
    notes(s, "Memory: ERA5 via Open-Meteo. Exposure: FAO-56 on Serbia’s real crops. Plan: Grok explains the measured numbers — it does not generate them.")

    # 8. Close
    s = blank(prs)
    eyebrow_box(s, Inches(0.8), Inches(1.55), Inches(11.7), "Measured  ·  not a forecast", PP_ALIGN.CENTER)
    box, tf = textbox(s, Inches(0.8), Inches(1.95), Inches(11.7), Inches(1.4))
    para(tf, "A decision.", 54, INK, SERIF, True, PP_ALIGN.CENTER, first=True)
    pins = [
        ("Subotica", "MASK", True),
        ("Novi Sad", "MASK", True),
        ("Zrenjanin", "URGENT", True),
        ("Belgrade", "FLAT", False),
        ("Kraljevo", "STABLE", False),
        ("Niš", "BETTER", False),
    ]
    pw, ph, pg = Inches(1.7), Inches(1.55), Inches(0.18)
    total = 6 * pw + 5 * pg
    x = (SLIDE_W - total) / 2
    py = Inches(4.05)
    for name, tag, hot in pins:
        fill = PANEL_DANGER if hot else PANEL
        card(s, x, py, pw, ph, fill, radius=0.18)
        dot_c = DANGER if hot else ACCENT
        oval(s, x + (pw - Inches(0.16)) / 2, py + Inches(0.22), Inches(0.16), Inches(0.16), dot_c)
        box, tf = textbox(s, x + Inches(0.06), py + Inches(0.48), pw - Inches(0.12), Inches(0.9))
        para(tf, name, 12, INK, SANS, True, PP_ALIGN.CENTER, first=True)
        para(tf, tag, 10, MUTED, SANS, True, PP_ALIGN.CENTER, space_before=2, tracking=8)
        x += pw + pg
    notes(s, "The pitch is not AI for farms. Annual rain can hide a worse summer, and the right response is different in Novi Sad, Zrenjanin, and Niš. Not a forecast. A decision.")

    prs.save(OUT)
    OUT_DOCS.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_DOCS)
    print(f"Wrote {OUT}")
    print(f"Wrote {OUT_DOCS}")


if __name__ == "__main__":
    build()
